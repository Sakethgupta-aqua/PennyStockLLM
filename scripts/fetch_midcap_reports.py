"""
Download the Nifty Midcap 100 broker reports listed in the Trendlyne spreadsheet.

The spreadsheet holds 297 rows - the latest 3 reports for each of 99 companies - and each
row's "PDF link" cell carries a real Trendlyne URL in its hyperlink target (the visible
text is just "Open PDF"). This reads those targets, lays out one folder per company, and
fetches the PDFs.

WHY A COOKIE IS REQUIRED. Trendlyne serves these PDFs only to a logged-in session. Fetched
without one, every URL returns HTTP 200 with `text/html` and the page title "Trendlyne
login" - a 60 KB login page, not a report. That is why this script verifies the bytes and
refuses to save an HTML file under a .pdf name: 297 login pages named *.pdf would poison
the corpus far more expensively than a failed download.

    # 1. log in to trendlyne.com in your browser
    # 2. export cookies (see --help-cookies) to cookies.txt
    python scripts/fetch_midcap_reports.py --layout            # folders + manifests only
    python scripts/fetch_midcap_reports.py --cookies cookies.txt --dry-run
    python scripts/fetch_midcap_reports.py --cookies cookies.txt

Resumable and safe to re-run: a report already on disk is skipped, so an interrupted run
picks up where it stopped.

THE QUOTA. Trendlyne counts every report open against the account - 70 a day (reset at
midnight IST), 300 a month (reset on the 1st) - retries included, and past it every link
redirects to /member/daily-runout/, whose own text says the cap is there to stop crawlers and
downloaders. So a run opens at most DAILY_OPENS reports, stops at the first runout page, and
spends its opens on the most useful reports first. 297 reports therefore take several days.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import time
from collections import Counter
from dataclasses import asdict, dataclass
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from pennystock.data.trendlyne_session import COOKIE_HELP, UA, load_cookies  # noqa: E402

XLSX = Path.home() / "Downloads" / "Nifty_Midcap100_Research_Reports.xlsx"
OUT = ROOT / "Broker_Reports"
SHEET = "Reports"

# Trendlyne rate-limits and this is someone else's server: one request at a time, with a
# pause. 297 reports at 1.5s is under 8 minutes, which is not worth risking a block over.
DELAY_S = 1.5
TIMEOUT_S = 60
MIN_PDF_BYTES = 10_000          # a real report is 300KB-2MB; anything tiny is an error page
DAILY_OPENS = 70                # Trendlyne's daily research-report allowance for this account
RUNOUT = "daily-runout"         # where every report link goes once the allowance is spent
# Brokers whose report links lead to a host that no longer resolves (gcc.geojit.net, measured
# 2026-10-02): opening one spends an open on a request that cannot succeed.
DEAD_HOSTS = ("geojit",)



@dataclass
class Row:
    """One spreadsheet row: a single broker report on a single company."""
    batch: str | None
    idx: int | None
    company: str | None
    ticker: str | None
    report_no: int | None
    report_date: str | None
    broker: str | None
    rating: str | None
    target: float | None
    change_vs_prev: str | None
    headline: str | None
    filename: str | None
    url: str | None


def _safe(name: str) -> str:
    """A folder name that survives every filesystem, without losing the ticker."""
    return re.sub(r"[^A-Za-z0-9._-]+", "_", (name or "UNKNOWN").strip()).strip("_") or "UNKNOWN"


def read_rows(xlsx: Path = XLSX) -> list[Row]:
    """Pull the 297 rows out of the sheet, URLs included.

    The URL lives in the cell's HYPERLINK TARGET, not its text - `values_only=True` would
    hand back the string "Open PDF" 297 times and silently lose every link.
    """
    import openpyxl

    wb = openpyxl.load_workbook(xlsx)
    if SHEET not in wb.sheetnames:
        raise SystemExit(f"{xlsx} has no '{SHEET}' sheet (found {wb.sheetnames})")
    ws = wb[SHEET]

    def cell(c):
        v = c.value
        return v.date().isoformat() if hasattr(v, "date") else v

    rows: list[Row] = []
    for r in ws.iter_rows(min_row=2):
        if not any(c.value for c in r):
            continue
        link = r[12].hyperlink.target if len(r) > 12 and r[12].hyperlink else None
        rows.append(Row(batch=cell(r[0]), idx=cell(r[1]), company=cell(r[2]),
                        ticker=cell(r[3]), report_no=cell(r[4]), report_date=cell(r[5]),
                        broker=cell(r[6]), rating=cell(r[7]), target=cell(r[8]),
                        change_vs_prev=cell(r[9]), headline=cell(r[10]),
                        filename=cell(r[11]), url=link))
    return rows


def target_path(row: Row) -> Path:
    """Where this report belongs: Broker_Reports/<TICKER>/<suggested filename>.

    The folder is the TICKER, not the display name, because the ticker is what every other
    table in this project keys on - stock_metrics.symbol, predictions.symbol, the resolver.
    A folder called "360ONE" joins to those; one called "360 One Wam" does not.
    """
    fname = _safe(row.filename or f"{row.ticker}_{row.report_no}.pdf")
    if not fname.lower().endswith(".pdf"):
        fname += ".pdf"
    return OUT / _safe(row.ticker) / fname


def write_layout(rows: list[Row]) -> tuple[int, int]:
    """Create one folder per company, each with a manifest describing its reports.

    The manifest is written whether or not the PDFs arrive. It is the part of this job that
    does not depend on being logged in, and it is independently useful: it records what
    exists, dated, with the broker's rating and target, so the set is auditable even while
    the bytes are missing.
    """
    by_ticker: dict[str, list[Row]] = {}
    for r in rows:
        by_ticker.setdefault(_safe(r.ticker), []).append(r)

    cols = ["report_no", "report_date", "broker", "rating", "target", "change_vs_prev",
            "headline", "filename", "url", "downloaded"]
    for ticker, group in by_ticker.items():
        folder = OUT / ticker
        folder.mkdir(parents=True, exist_ok=True)
        with (folder / "manifest.csv").open("w", newline="", encoding="utf-8") as fh:
            w = csv.DictWriter(fh, fieldnames=["company", "ticker", *cols])
            w.writeheader()
            for r in sorted(group, key=lambda x: x.report_no or 0):
                w.writerow({"company": r.company, "ticker": r.ticker,
                            **{c: getattr(r, c, None) for c in cols if c != "downloaded"},
                            "downloaded": "Y" if target_path(r).exists() else ""})

    (OUT / "midcap100_manifest.json").write_text(
        json.dumps([asdict(r) | {"local_path": str(target_path(r).relative_to(ROOT)),
                                 "downloaded": target_path(r).exists()} for r in rows],
                   indent=2), encoding="utf-8")
    return len(by_ticker), len(rows)


def _dead_host(row: Row) -> bool:
    return any(h in (row.broker or "").lower() for h in DEAD_HOSTS)


def fetch_order(rows: list[Row]) -> list[Row]:
    """The reports still to fetch, most useful first.

    One company's newest report is worth more than another company's third, so the queue runs in
    rounds: each company's next missing report - newest reports first - before anyone's report
    after it. Reports on a dead host are left out.
    """
    have = Counter(r.ticker for r in rows if target_path(r).exists())
    queued: Counter = Counter()
    keyed = []
    for r in sorted(rows, key=lambda r: r.report_no or 0):
        if target_path(r).exists() or _dead_host(r):
            continue
        keyed.append((have[r.ticker] + queued[r.ticker], r))
        queued[r.ticker] += 1
    keyed.sort(key=lambda kr: str(kr[1].report_date or ""), reverse=True)
    keyed.sort(key=lambda kr: kr[0])                  # stable: newest first within each round
    return [r for _, r in keyed]


def fetch(rows: list[Row], cookies: Path | None, cookie_header: str | None,
          dry_run: bool = False, limit: int | None = DAILY_OPENS,
          delay: float = DELAY_S, session=None) -> dict:
    """Fetch reports not already on disk, verifying each one is really a PDF - at most `limit`
    opens, and none after Trendlyne's daily-runout page."""
    import requests

    session = session or requests.Session()
    session.headers.update({"User-Agent": UA, "Referer": "https://trendlyne.com/"})
    load_cookies(session, cookies, cookie_header)

    tally = {"saved": 0, "skipped": 0, "login_wall": 0, "not_pdf": 0, "failed": 0, "no_url": 0,
             "dead_host": 0, "daily_limit": 0}
    queue = fetch_order(rows)
    todo = queue[:limit] if limit else queue
    tally["skipped"] = sum(1 for r in rows if target_path(r).exists())
    tally["dead_host"] = sum(1 for r in rows if _dead_host(r) and not target_path(r).exists())

    print(f"{len(todo)} to fetch now, {len(queue) - len(todo)} left for later runs, "
          f"{tally['skipped']} already on disk, {tally['dead_host']} on a dead host (skipped)\n")
    for i, row in enumerate(todo, 1):
        dest = target_path(row)
        label = f"[{i}/{len(todo)}] {row.ticker} #{row.report_no} {row.broker}"
        if not row.url:
            tally["no_url"] += 1
            print(f"{label}: no URL in the sheet")
            continue
        if dry_run:
            print(f"{label}: would GET {row.url} -> {dest.relative_to(ROOT)}")
            continue

        try:
            resp = session.get(row.url, timeout=TIMEOUT_S, allow_redirects=True)
        except Exception as e:
            tally["failed"] += 1
            print(f"{label}: request failed - {type(e).__name__}: {e}")
            continue

        body, ctype = resp.content, resp.headers.get("Content-Type", "")
        # Verify the BYTES, never the status code. Trendlyne answers an unauthenticated
        # request with HTTP 200 and a login page, so status alone would call that a success.
        if body[:5] != b"%PDF-":
            if RUNOUT in str(getattr(resp, "url", "")) or RUNOUT.encode() in body[:80_000]:
                # every further open would land here too, and still count against the account
                tally["daily_limit"] += 1
                print(f"{label}: Trendlyne's daily report limit is used up - stopping. "
                      "It resets at midnight IST; run again then.")
                break
            if b"Trendlyne login" in body[:80_000] or "text/html" in ctype:
                tally["login_wall"] += 1
                print(f"{label}: LOGIN PAGE returned - cookie missing or expired")
                if tally["login_wall"] >= 3 and tally["saved"] == 0:
                    print("\nStopping: the first 3 fetches all hit the login wall.\n"
                          + COOKIE_HELP)
                    break
            else:
                tally["not_pdf"] += 1
                print(f"{label}: not a PDF ({ctype}, {len(body)}B)")
            continue
        if len(body) < MIN_PDF_BYTES:
            tally["not_pdf"] += 1
            print(f"{label}: PDF too small to be a report ({len(body)}B)")
            continue

        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_bytes(body)
        tally["saved"] += 1
        print(f"{label}: saved {len(body):,}B -> {dest.relative_to(ROOT)}")
        time.sleep(delay)

    return tally


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", type=Path, default=XLSX)
    ap.add_argument("--cookies", type=Path, help="Netscape cookies.txt from a logged-in browser")
    ap.add_argument("--cookie-header", help='raw header: "sessionid=...; csrftoken=..."')
    ap.add_argument("--layout", action="store_true", help="only create folders and manifests")
    ap.add_argument("--dry-run", action="store_true", help="show what would be fetched")
    ap.add_argument("--limit", type=int, default=DAILY_OPENS,
                    help=f"open at most N reports (default {DAILY_OPENS}, Trendlyne's daily "
                         "allowance; use 3 to test the cookie)")
    ap.add_argument("--delay", type=float, default=DELAY_S)
    ap.add_argument("--help-cookies", action="store_true", help="explain how to export cookies")
    args = ap.parse_args()

    if args.help_cookies:
        print(COOKIE_HELP)
        return 0

    rows = read_rows(args.xlsx)
    companies, total = write_layout(rows)
    have = sum(1 for r in rows if target_path(r).exists())
    print(f"{total} reports across {companies} companies; {have} already on disk")
    print(f"folders + manifest.csv written under {OUT.relative_to(ROOT)}/")
    print(f"index: {(OUT / 'midcap100_manifest.json').relative_to(ROOT)}\n")

    if args.layout:
        return 0
    if not (args.cookies or args.cookie_header or args.dry_run):
        print("No cookie given, and Trendlyne serves these PDFs only to a logged-in "
              "session.\nRe-run with --cookies cookies.txt (see --help-cookies).")
        return 2

    tally = fetch(rows, args.cookies, args.cookie_header, args.dry_run, args.limit, args.delay)
    print("\n" + "  ".join(f"{k}={v}" for k, v in tally.items()))
    if not args.dry_run:
        write_layout(rows)                     # refresh the `downloaded` column
    return 1 if tally["login_wall"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
